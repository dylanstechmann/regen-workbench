#!/usr/bin/env python3
"""Regenerate compact descriptive tables from the cited public workbooks."""
from __future__ import annotations

import argparse
import csv
import hashlib
import itertools
import json
import math
import os
import re
import statistics
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
DEFAULT_SNAPSHOTS = Path(os.environ.get("REGEN_DATA", ROOT / "data")) / "research-desk" / "source_snapshots"
DEFAULT_OUTPUT = HERE / "derived"
SOURCE_MANIFEST = HERE / "source_manifest.json"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def checked_sources(directory: Path) -> dict[str, dict]:
    manifest = json.loads(SOURCE_MANIFEST.read_text(encoding="utf-8"))
    result = {}
    for source in manifest["snapshots"]:
        path = directory / source["file"]
        if not path.is_file():
            raise FileNotFoundError(f"missing source snapshot: {path}")
        digest = sha256_file(path)
        if digest != source["sha256"] or path.stat().st_size != source["bytes"]:
            raise ValueError(f"source snapshot failed pinned size/SHA-256 check: {source['file']}")
        result[source["file"]] = {**source, "path": path, "verified_sha256": digest}
    return result


def workbook(path: Path):
    return load_workbook(path, read_only=True, data_only=True)


def body_map(sources: dict[str, dict]) -> dict:
    path = sources["body_map_somatic_mutations_supplementary_table3.xlsx"]["path"]
    book = workbook(path)
    sheet = book["Sheet1"]
    rows = sheet.iter_rows(values_only=True)
    next(rows)
    next(rows)
    header = next(rows)
    index = {name: i for i, name in enumerate(header)}
    if not {"sampleID", "impact", "gene"}.issubset(index):
        raise ValueError("unexpected body-map Table 3 schema")
    total = 0
    sample_ids = set()
    liver_impacts = Counter()
    liver_rows = 0
    date_coerced_gene_rows = 0
    for row in rows:
        if not row or row[index["sampleID"]] is None:
            continue
        total += 1
        sample_ids.add(str(row[index["sampleID"]]))
        if isinstance(row[index["gene"]], (date, datetime)):
            date_coerced_gene_rows += 1
        sample = str(row[index["sampleID"]])
        if re.match(r"^PN\d+L-", sample):
            liver_rows += 1
            liver_impacts[str(row[index["impact"]] or "unknown")] += 1
    book.close()

    path = sources["body_map_somatic_mutations_supplementary_table10.xlsx"]["path"]
    book = workbook(path)
    sheet = book["p-value"]
    rows = sheet.iter_rows(values_only=True)
    next(rows)
    header = next(rows)
    organs = [str(value) for value in header[1:] if value is not None]
    liver_index = next((i + 1 for i, value in enumerate(header[1:]) if value and "liver" in str(value).lower()), None)
    if liver_index is None:
        raise ValueError("body-map Table 10 has no liver column")
    tests = []
    liver_test_indexes = []
    for row in rows:
        if row and row[0] is not None:
            for organ_index, p_value in enumerate(row[1:]):
                if p_value is not None and math.isfinite(float(p_value)):
                    tests.append((str(row[0]), organs[organ_index], float(p_value)))
                    if organ_index + 1 == liver_index:
                        liver_test_indexes.append(len(tests) - 1)
    book.close()
    ordered = sorted((p, i) for i, (_, _, p) in enumerate(tests))
    adjusted = [1.0] * len(tests)
    running = 1.0
    for rank in range(len(ordered), 0, -1):
        p, original_index = ordered[rank - 1]
        running = min(running, p * len(tests) / rank)
        adjusted[original_index] = min(1.0, running)
    liver = [{"gene": tests[i][0], "p_value": tests[i][2], "bh_q_value": adjusted[i]} for i in liver_test_indexes]
    liver.sort(key=lambda row: (row["p_value"], row["gene"]))
    return {
        "coding_variant_rows": total,
        "sample_count_with_at_least_one_row": len(sample_ids),
        "liver_coding_variant_rows": liver_rows,
        "liver_impact_counts": dict(sorted(liver_impacts.items())),
        "date_coerced_gene_cells": date_coerced_gene_rows,
        "driver_test_count": len(tests),
        "liver_driver_enrichment": liver,
        "interpretation": "Rows and one-sided enrichment tests are not per-cell burdens, causal targets, or correction results.",
    }


def design_matrix(records: list[dict], predictors: tuple[str, ...]) -> np.ndarray:
    columns = [np.ones(len(records), dtype=float)]
    for predictor in predictors:
        values = [row[predictor] for row in records]
        if predictor == "age":
            columns.append(np.asarray(values, dtype=float))
        else:
            levels = sorted({str(value) for value in values})
            for level in levels[1:]:
                columns.append(np.asarray([float(str(value) == level) for value in values]))
    return np.column_stack(columns)


def r_squared(records: list[dict], predictors: tuple[str, ...], response: str) -> float:
    y = np.asarray([float(row[response]) for row in records], dtype=float)
    design = design_matrix(records, predictors)
    coefficients, _, _, _ = np.linalg.lstsq(design, y, rcond=None)
    residual = y - design @ coefficients
    total = y - float(y.mean())
    return float(1 - (residual @ residual) / (total @ total)) if float(total @ total) else 0.0


def log_count_age_fit(rows: list[dict]) -> tuple[float, float]:
    ages = np.asarray([float(row["age"]) for row in rows], dtype=float) - 65.0
    log_calls = np.log(np.asarray([int(row["calls"]) for row in rows], dtype=float))
    centered_age = ages - ages.mean()
    denominator = float(centered_age @ centered_age)
    slope = float(centered_age @ (log_calls - log_calls.mean()) / denominator) if denominator else 0.0
    return float(log_calls.mean() - slope * ages.mean()), slope


def bootstrap_age_predictions(rows: list[dict], rng: np.random.Generator, samples=1000) -> np.ndarray:
    ages = np.asarray([float(row["age"]) for row in rows], dtype=float) - 65.0
    log_calls = np.log(np.asarray([int(row["calls"]) for row in rows], dtype=float))
    indices = rng.integers(0, len(rows), size=(samples, len(rows)))
    x = ages[indices]
    y = log_calls[indices]
    x_centered = x - x.mean(axis=1, keepdims=True)
    y_centered = y - y.mean(axis=1, keepdims=True)
    denominator = np.sum(x_centered * x_centered, axis=1)
    numerator = np.sum(x_centered * y_centered, axis=1)
    slopes = np.divide(numerator, denominator, out=np.zeros_like(numerator), where=denominator != 0)
    intercepts = y.mean(axis=1) - slopes * x.mean(axis=1)
    return np.exp(intercepts)


def skin_study(sources: dict[str, dict]) -> tuple[dict, list[dict]]:
    book = workbook(sources["skin_supplementary_datasets_s1_s2.xlsx"]["path"])
    variants_sheet = book["Dataset_S1"]
    variant_rows = variants_sheet.iter_rows(values_only=True)
    for row in variant_rows:
        if row and row[0] == "sampleID":
            headers = list(row)
            break
    else:
        raise ValueError("skin Dataset S1 header not found")
    col = {name: i for i, name in enumerate(headers)}
    call_counts = Counter()
    cancer_gene_donors = set()
    annotation_counts = Counter()
    damaging_terms = ("missense", "nonsense", "frameshift", "frame_shift", "splice", "in-frame", "inframe", "stop_gain", "stop gain")
    for row in variant_rows:
        if not row or row[col["sampleID"]] is None:
            continue
        sample = str(row[col["sampleID"]])
        call_counts[sample] += 1
        annotation = str(row[col["annotation"]] or "").lower()
        annotation_counts[annotation or "unknown"] += 1
        gene = str(row[col["gene_name"]] or "").upper()
        if gene in {"TP53", "NOTCH1", "FAT1"} and any(term in annotation for term in damaging_terms):
            cancer_gene_donors.add(sample)

    meta_sheet = book["Dataset_S2"]
    meta_rows = meta_sheet.iter_rows(values_only=True)
    for row in meta_rows:
        if row and row[0] == "sampleID":
            meta_headers = list(row)
            break
    else:
        raise ValueError("skin Dataset S2 header not found")
    meta_col = {name: i for i, name in enumerate(meta_headers)}
    records = []
    for row in meta_rows:
        if not row or row[meta_col["sampleID"]] is None:
            continue
        record = {key: row[index] for key, index in meta_col.items()}
        record["calls"] = int(call_counts.get(str(record["sampleID"]), 0))
        records.append(record)
    book.close()
    ids = [str(row["sampleID"]) for row in records]
    if len(ids) != len(set(ids)) or set(ids) != set(call_counts):
        raise ValueError("skin S1/S2 join is not one-to-one")

    photo = defaultdict(list)
    for row in records:
        photo[str(row["skin_phototype"]) if row["skin_phototype"] is not None else "missing"].append(row)
    photo_summary = []
    bootstrap_rng = np.random.default_rng(12345)
    for group in sorted(photo):
        group_rows = photo[group]
        ages = [float(row["age"]) for row in group_rows]
        calls = [int(row["calls"]) for row in group_rows]
        estimate = {
            "phototype": group,
            "donors": len(group_rows),
            "mean_age_years": statistics.mean(ages),
            "median_calls": statistics.median(calls),
            "age_multiplier_per_decade": None,
            "predicted_calls_at_65": None,
            "bootstrap_95_percent_interval_at_65": None,
            "age_model_status": "insufficient phototype/age observations" if len(group_rows) < 2 or len(set(ages)) < 2 else "computed",
        }
        if estimate["age_model_status"] == "computed":
            intercept, slope = log_count_age_fit(group_rows)
            predictions = bootstrap_age_predictions(group_rows, bootstrap_rng)
            estimate.update({
                "age_multiplier_per_decade": float(np.exp(slope * 10)),
                "predicted_calls_at_65": float(np.exp(intercept)),
                "bootstrap_95_percent_interval_at_65": [float(x) for x in np.quantile(predictions, [0.025, 0.975])],
            })
        photo_summary.append(estimate)

    predictors = ("age", "skin_phototype", "sex", "UV_exposure_tissue", "sun_damage_tissue", "sun_history", "MC1R_genotype")
    complete = [row for row in records if all(row.get(field) is not None for field in predictors)]
    complete = [row for row in complete if int(row["calls"]) > 0]
    for row in complete:
        row["log_calls"] = math.log(int(row["calls"]))
    full_r2 = r_squared(complete, predictors, "log_calls")
    parameter_count = design_matrix(complete, predictors).shape[1] - 1
    adjusted_r2 = 1 - (1 - full_r2) * (len(complete) - 1) / (len(complete) - parameter_count - 1)
    lmg = {}
    n_predictors = len(predictors)
    r2_cache = {(): r_squared(complete, (), "log_calls")}
    for size in range(1, n_predictors + 1):
        for subset in itertools.combinations(predictors, size):
            r2_cache[subset] = r_squared(complete, subset, "log_calls")
    denominator = math.factorial(n_predictors)
    for predictor in predictors:
        others = [item for item in predictors if item != predictor]
        contribution = 0.0
        for size in range(n_predictors):
            weight = math.factorial(size) * math.factorial(n_predictors - size - 1) / denominator
            for subset in itertools.combinations(others, size):
                before = tuple(sorted(subset, key=predictors.index))
                after = tuple(sorted((*subset, predictor), key=predictors.index))
                contribution += weight * (r2_cache[after] - r2_cache[before])
        lmg[predictor] = contribution / full_r2 * 100 if full_r2 else 0.0
    summary = {
        "variant_rows": sum(call_counts.values()),
        "donors_joined": len(records),
        "mean_calls_per_donor": statistics.mean(call_counts.values()),
        "median_calls_per_donor": statistics.median(call_counts.values()),
        "min_calls_per_donor": min(call_counts.values()),
        "max_calls_per_donor": max(call_counts.values()),
        "donors_with_protein_altering_TP53_NOTCH1_or_FAT1_call": len(cancer_gene_donors),
        "protein_altering_annotation_rule": list(damaging_terms),
        "annotation_counts": dict(sorted(annotation_counts.items())),
        "complete_case_n": len(complete),
        "full_model_r_squared": full_r2,
        "full_model_adjusted_r_squared": adjusted_r2,
        "lmg_share_percent": lmg,
        "by_phototype": photo_summary,
        "interpretation": "Targeted-panel clone counts are not whole-genome burden, causal effects, edit targets, or evidence that correction reverses skin aging.",
    }
    return summary, photo_summary


def nanoseq_summary(path: Path, tissue: str) -> tuple[dict, list[dict]]:
    book = workbook(path)
    sheet = book["Sheet1"]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    col = {name: i for i, name in enumerate(headers)}
    if not {"type", "duplex_vaf"}.issubset(col):
        raise ValueError(f"unexpected NanoSeq table schema: {path.name}")
    types = Counter()
    vafs = []
    total = 0
    for row in rows:
        if not row or row[col["type"]] is None:
            continue
        total += 1
        types[str(row[col["type"]])] += 1
        if row[col["duplex_vaf"]] is not None:
            value = float(row[col["duplex_vaf"]])
            if math.isfinite(value):
                vafs.append(value)
    book.close()
    values = np.asarray(vafs, dtype=float)
    quantiles = np.quantile(values, [0.05, 0.5, 0.95], method="linear")
    result = {"tissue": tissue, "mutation_rows": total, "rows_with_duplex_vaf": len(vafs), "type_counts": dict(sorted(types.items())),
              "duplex_vaf_p05": float(quantiles[0]), "duplex_vaf_median": float(quantiles[1]), "duplex_vaf_p95": float(quantiles[2]),
              "interpretation": "Mutation-row distribution; no participant identifiers or person-level prevalence denominator."}
    type_rows = [{"tissue": tissue, "variant_type": kind, "rows": count} for kind, count in sorted(types.items())]
    return result, type_rows


def site_selection(path: Path) -> dict:
    book = workbook(path)
    result = {}
    for sheet_name in ("sitednds_oral", "sitednds_blood"):
        sheet = book[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        col = {name: i for i, name in enumerate(header)}
        genes = set()
        count = 0
        for row in rows:
            if row and row[col["gene"]] is not None:
                count += 1
                genes.add(str(row[col["gene"]]))
        result[sheet_name.removeprefix("sitednds_")] = {"selected_site_rows": count, "distinct_genes": len(genes)}
    book.close()
    return result


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def atomic_text(path: Path, text: str) -> None:
    temporary = path.with_name(path.name + f".{os.getpid()}.tmp")
    temporary.write_text(text, encoding="utf-8")
    temporary.replace(path)


def run(snapshots: Path, output: Path) -> dict:
    sources = checked_sources(snapshots)
    body = body_map(sources)
    skin, skin_groups = skin_study(sources)
    oral, oral_types = nanoseq_summary(sources["nanoseq_oral_mutations_supplementary_table8.xlsx"]["path"], "oral")
    blood, blood_types = nanoseq_summary(sources["nanoseq_blood_mutations_supplementary_table9.xlsx"]["path"], "blood")
    result = {
        "schema_version": 1,
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "source_files": [{k: v for k, v in item.items() if k != "path"} for item in sources.values()],
        "body_map": body,
        "skin": skin,
        "nanoseq": {"oral": oral, "blood": blood, "site_selection": site_selection(sources["nanoseq_site_selected_replacements_supplementary_table4.xlsx"]["path"]),
                    "interpretation": "Site-level selection and mutation-row VAF summaries are not participant-level outcomes or beneficial repair targets."},
        "overall_interpretation": "Descriptive reconstruction of public supplementary tables. No source study tested broad correction of ordinary age-acquired mutations as a rejuvenation intervention.",
    }
    output.mkdir(parents=True, exist_ok=True)
    rows = body["liver_driver_enrichment"]
    write_csv(output / "liver_driver_enrichment.csv", ["gene", "p_value", "bh_q_value"], rows)
    write_csv(output / "bodymap_liver_impact.csv", ["impact", "rows"], [{"impact": key, "rows": value} for key, value in body["liver_impact_counts"].items()])
    write_csv(output / "skin_phototype_summary.csv", ["phototype", "donors", "mean_age_years", "median_calls", "age_model_status", "age_multiplier_per_decade", "predicted_calls_at_65", "bootstrap_95_percent_interval_at_65"],
              [{**row, "bootstrap_95_percent_interval_at_65": json.dumps(row["bootstrap_95_percent_interval_at_65"])} for row in skin_groups])
    write_csv(output / "nanoseq_variant_types.csv", ["tissue", "variant_type", "rows"], oral_types + blood_types)
    atomic_text(output / "summary.json", json.dumps(result, ensure_ascii=True, indent=2, allow_nan=False) + "\n")
    brief = ["# Mutation-repair pilot: reproducible descriptive tables", "",
             "These are aggregate summaries of public supplementary workbooks. They are not per-cell mutation burdens, causal repair targets, or evidence of rejuvenation.", "",
             f"- Body-map coding calls: {body['coding_variant_rows']:,} rows in {body['sample_count_with_at_least_one_row']:,} samples; {body['liver_coding_variant_rows']:,} liver rows.",
             f"- Skin panel: {skin['variant_rows']:,} calls joined one-to-one to {skin['donors_joined']} donor records; median {skin['median_calls_per_donor']} calls per donor.",
             f"- NanoSeq mutation rows: oral {oral['mutation_rows']:,}; blood {blood['mutation_rows']:,}. No participant identifiers are present in these public tables.",
             f"- Body-map driver-enrichment tests: {body['driver_test_count']} gene-tissue tests; BH adjustment is computed across the full table.",
             "", "The skin model uses explicit categorical dummy variables and ordinary least squares. Bootstrap intervals resample donors within each phototype in sorted order using one NumPy generator with fixed seed 12345. Differences from the manuscript can arise from spreadsheet coercion, formula choices or study-specific analysis code.",
             "", "A mutation or expanding clone is not automatically harmful. These tables contain no intervention that corrects ordinary age-acquired variants and restores tissue function.", ""]
    atomic_text(output / "SUMMARY.md", "\n".join(brief))
    output_files = [p for p in output.iterdir() if p.is_file() and p.name != "manifest.json"]
    manifest = {"schema_version": 1, "script_sha256": sha256_file(Path(__file__)),
                "source_manifest_sha256": sha256_file(SOURCE_MANIFEST),
                "inputs": {name: {"sha256": item["sha256"], "bytes": item["bytes"], "url": item["url"]} for name, item in sources.items()},
                "outputs": {p.name: {"sha256": sha256_file(p), "bytes": p.stat().st_size} for p in sorted(output_files)}}
    atomic_text(output / "manifest.json", json.dumps(manifest, ensure_ascii=True, indent=2) + "\n")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshots", type=Path, default=DEFAULT_SNAPSHOTS)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = run(args.snapshots, args.out)
    print(json.dumps({"output": str(args.out), "body_map_rows": result["body_map"]["coding_variant_rows"],
                      "skin_donors": result["skin"]["donors_joined"], "oral_mutation_rows": result["nanoseq"]["oral"]["mutation_rows"],
                      "blood_mutation_rows": result["nanoseq"]["blood"]["mutation_rows"]}, indent=2))


if __name__ == "__main__":
    main()
