"""Read-only reanalysis of Quintard 2024 Fig7 paired insulin observations.

Descriptive spheroid-level outcomes; donor/batch inference is unavailable.
Source workbook is never edited. Requires openpyxl only for the source intake.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import statistics

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]


def text_hash(path):
    return hashlib.sha256(path.read_text(encoding="utf-8").encode()).hexdigest()


def measured(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def extract_pairs(sheet, plan):
    if sheet["A1"].value != "Figure 7c" or sheet["A38"].value != "Figure 7d" or sheet["K2"].value != "(pg/ml)":
        raise ValueError("Source figure/units layout changed")
    output = []
    for group in plan["groups"]:
        col = group["low_column"]
        if (sheet.cell(2, col).value != group["source_label"]
                or sheet.cell(3, col).value != "Low Glucose"
                or sheet.cell(3, col + 1).value != "High Glucose"
                or sheet.cell(40, group["si_column"]).value != group["source_label"]):
            raise ValueError("Source group headers changed")
        ended = False
        records = []
        for row in range(4, 26):
            low_cell, high_cell = sheet.cell(row, col), sheet.cell(row, col + 1)
            low, high = low_cell.value, high_cell.value
            source_si = sheet.cell(row + 37, group["si_column"])
            if low is None and high is None:
                if source_si.value is not None:
                    raise ValueError("Source SI lacks its paired observations")
                ended = True
                continue
            if ended or not measured(low) or not measured(high) or low <= 0 or high < 0:
                raise ValueError("Invalid, interrupted or incomplete low/high pair")
            ratio = high / low
            if not measured(source_si.value) or abs(ratio - source_si.value) > 1e-10:
                raise ValueError("Recomputed high/low ratio differs from source SI")
            records.append({"group": group["id"], "source_label": group["source_label"],
                "source_low_cell": f"Figure 7!{low_cell.coordinate}",
                "source_high_cell": f"Figure 7!{high_cell.coordinate}",
                "source_si_cell": f"Figure 7!{source_si.coordinate}",
                "reported_point_id": f'{group["id"]}_row{row}',
                "low_insulin_pg_per_ml": float(low), "high_insulin_pg_per_ml": float(high),
                "stimulation_index_high_over_low": ratio, "source_stimulation_index": source_si.value,
                "ratio_reconstruction_error": ratio - source_si.value,
                "donor_id": None, "manufacturing_batch_id": None,
                "unit": "reported pancreatic islet spheroid observation"})
        if len(records) != group["expected_observations"]:
            raise ValueError("Observed group size differs from pinned source layout")
        output.extend(records)
    return output


def summarize(rows):
    result = []
    for group in dict.fromkeys(row["group"] for row in rows):
        selected = [row for row in rows if row["group"] == group]
        ratios = [row["stimulation_index_high_over_low"] for row in selected]
        low = statistics.mean(row["low_insulin_pg_per_ml"] for row in selected)
        high = statistics.mean(row["high_insulin_pg_per_ml"] for row in selected)
        result.append({"group": group, "n_reported_spheroid_observations": len(selected),
            "mean_stimulation_index": statistics.mean(ratios),
            "median_stimulation_index": statistics.median(ratios),
            "sample_sd_stimulation_index": statistics.stdev(ratios) if len(ratios) > 1 else None,
            "minimum_stimulation_index": min(ratios), "maximum_stimulation_index": max(ratios),
            "mean_low_insulin_pg_per_ml": low, "mean_high_insulin_pg_per_ml": high,
            "ratio_of_group_mean_concentrations": high / low,
            "n_independent_human_donors": None, "n_independent_manufacturing_batches": None})
    return result


def compare(rows):
    groups = {row["group"]: [r for r in rows if r["group"] == row["group"]] for row in rows}
    treated = groups["flow_with_vascularization"]
    contrasts, omissions = [], []
    for group, controls in groups.items():
        if group == "flow_with_vascularization":
            continue
        a = statistics.mean(row["stimulation_index_high_over_low"] for row in treated)
        b = statistics.mean(row["stimulation_index_high_over_low"] for row in controls)
        for role, current in (("flow_with_vascularization", treated), (group, controls)):
            if len(current) < 2:
                raise ValueError("Omission comparison requires two observations per condition")
            for dropped in current:
                remaining = [row["stimulation_index_high_over_low"] for row in current if row is not dropped]
                effect = statistics.mean(remaining) - b if role == "flow_with_vascularization" else a - statistics.mean(remaining)
                omissions.append({"comparator": group, "omitted_from": role,
                    "omitted_point_id": dropped["reported_point_id"],
                    "mean_si_difference_after_single_observation_omission": effect})
        effects = [r["mean_si_difference_after_single_observation_omission"] for r in omissions if r["comparator"] == group]
        contrasts.append({"comparator": group, "mean_si_difference": a - b,
            "ratio_of_mean_stimulation_indices": a / b if b > 0 else None,
            "single_observation_omission_min_difference": min(effects),
            "single_observation_omission_max_difference": max(effects),
            "interpretation": "Descriptive assay observations; no donor/batch-independent causal effect or confidence interval"})
    means = {r["group"]: r["mean_stimulation_index"] for r in summarize(rows)}
    interaction = (means["flow_with_vascularization"] - means["flow_without_vascularization"]
                   - means["static_with_vascularization"] + means["static_without_vascularization"])
    return contrasts, omissions, interaction


def write_tsv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter="\t", lineterminator="\n")
        writer.writeheader(); writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--plan", type=Path, default=HERE / "analysis_plan.json")
    parser.add_argument("--workbook", type=Path, default=ROOT / "data/literature/organoid-perfusion-2026-10-19/Quintard_2024_source_data.xlsx")
    parser.add_argument("--out", type=Path, default=HERE / "derived")
    args = parser.parse_args()
    plan = json.loads(args.plan.read_text())
    digest = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    if digest != plan["source_workbook_sha256"]:
        raise ValueError("Source workbook hash changed")
    import openpyxl
    book = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = book["Figure 7"]
    rows = extract_pairs(sheet, plan)
    contrasts, omissions, interaction = compare(rows)
    args.out.mkdir(parents=True, exist_ok=True)
    outputs = {"paired_insulin_observations.tsv": rows, "condition_summary.tsv": summarize(rows),
               "condition_contrasts.tsv": contrasts, "single_observation_influence.tsv": omissions}
    for name, values in outputs.items(): write_tsv(args.out / name, values)
    context = {"source_statistic_scope": "Original workbook outputs transcribed; no confirmatory test recalculated",
        "figure_caption_test_description": "unpaired t test, two-tailed",
        "workbook_test_description": "ANOVA and Dunnett multiple comparisons",
        "source_anova_f": sheet["B70"].value, "source_anova_p": sheet["B71"].value,
        "source_variance_test": "Brown-Forsythe", "source_variance_test_p": sheet["B78"].value,
        "source_total_observations": sheet["B95"].value,
        "claim_limit": "No donor, lot or chip/batch mapping supplied in this workbook; reported spheroids are not independent human donors"}
    (args.out / "source_statistics_context.json").write_text(json.dumps(context, indent=2) + "\n", encoding="utf-8", newline="\n")
    book.close()
    if hashlib.sha256(args.workbook.read_bytes()).hexdigest() != digest:
        raise RuntimeError("Source workbook changed during read-only intake")
    result = {"source_doi": plan["source_doi"], "source_workbook_sha256": digest,
        "source_workbook_bytes": args.workbook.stat().st_size,
        "plan_sha256_lf_text": text_hash(args.plan), "script_sha256_lf_text": text_hash(Path(__file__)),
        "source_workbook_unchanged": True, "observations": len(rows), "condition_count": len(plan["groups"]),
        "single_observation_omission_rows": len(omissions),
        "max_source_si_reconstruction_error": max(abs(r["ratio_reconstruction_error"]) for r in rows),
        "descriptive_factorial_mean_contrast": interaction,
        "factorial_interpretation": "Flow+vascular-bed mean - flow-only mean - bed-only mean + static mean; no causal interaction inference",
        "output_sha256_lf_text": {name: text_hash(args.out / name) for name in list(outputs) + ["source_statistics_context.json"]},
        "limits": plan["limits"], "new_biological_experiment": False, "source_derived_license": "CC-BY-4.0"}
    (args.out / "analysis_receipt.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__": main()
